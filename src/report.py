"""One-click reports from a project: a design report of one variant, or a trade study
of several, as Word (.docx) or PDF.

Content comes from the project graph (inputs, results, what is stale) and the trust
layer (confidence, evidence and limits of every tool, validation references), so a
report always says how far each number can be trusted and where it came from.

    python report.py design ../cases/projects/trainer-study --variant baseline --format pdf
    python report.py trade ../cases/projects/trainer-study --variants baseline,heavy,cfrp --format docx
"""
import argparse
import datetime as dt
import json
import math
import subprocess
from pathlib import Path

import numpy as np

SRC = Path(__file__).resolve().parent
ROOT = SRC.parent
COLOR = {"green": "#2B7A4B", "amber": "#B26A00", "red": "#9C3D54", "info": "#56606C", None: "#56606C"}
FIGURES = {"drag": ["drag_buildup.png"], "loads": ["vn_diagram.png", "span_loads.png"],
           "performance": ["performance.png"], "polar": ["polar.png"]}


# ----------------------------------------------------------------------
# document model: a list of blocks, rendered to Word or PDF
# ----------------------------------------------------------------------
def H(level, text):
    return ("h", level, text)


def P(text, style=None):
    return ("p", text, style)


def BUL(items):
    return ("bullets", list(items))


def TAB(header, rows, colors=None, widths=None):
    """colors: {(row, col): level} for confidence / better-worse cells"""
    return ("table", header, rows, colors or {}, widths)


def IMG(path, caption=None, width=6.3):
    return ("img", str(path), caption, width)


def BREAK():
    return ("break",)


def _fmt(v, nd=4):
    if v is None:
        return "—"
    if isinstance(v, float):
        if not math.isfinite(v):
            return "—"
        return f"{v:.{nd}g}"
    return str(v)


def _git_version():
    try:
        r = subprocess.run(["git", "-C", str(ROOT), "rev-parse", "--short", "HEAD"], capture_output=True,
                           text=True, timeout=5)
        return r.stdout.strip() or "unknown"
    except Exception:                                    # noqa: BLE001
        return "unknown"


# ----------------------------------------------------------------------
# figures
# ----------------------------------------------------------------------
def planform_figure(ac, path):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    import aircraft as A
    fig, ax = plt.subplots(figsize=(6.4, 4.2))
    for s in ac.surfaces:
        for kind, q in A.surface_quads(s):
            loop = np.r_[q, q[:1]]
            if kind == "surface":
                ax.fill(loop[:, 1], -loop[:, 0], color="#1D4E89", alpha=0.08)
                ax.plot(loop[:, 1], -loop[:, 0], color="#1D4E89", lw=1)
            else:
                ax.fill(loop[:, 1], -loop[:, 0], color="#C2571A", alpha=0.25, lw=0)
    fu = ac.fuselage
    if fu.length > 0:
        xs = fu.x + np.linspace(0, fu.length, 60)
        r = fu.radius(xs)
        ax.plot(np.r_[r, -r[::-1]], -np.r_[xs, xs[::-1]], color="#56606C", lw=1)
    ax.plot([ac.cg[1]], [-ac.cg[0]], "D", color="#17202B", ms=5)
    ax.annotate("c.g.", (ac.cg[1], -ac.cg[0]), xytext=(6, 0), textcoords="offset points", fontsize=8)
    ax.set_aspect("equal")
    ax.set_xlabel("y (m)")
    ax.set_ylabel("−x (m)")
    ax.set_title(f"{ac.name}: top view", fontsize=10)
    ax.grid(alpha=0.25)
    fig.tight_layout()
    fig.savefig(path, dpi=150)
    plt.close(fig)
    return path


def structure_figure(d, path):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    s = d["structure"]
    y = np.array(s["y"])
    fig, ax = plt.subplots(1, 3, figsize=(9.5, 3.0))
    ax[0].plot(y, np.array(s["EI"]) / 1e6)
    ax[0].set_title("Bending stiffness EI (MN m²)", fontsize=9)
    ax[1].plot(y, np.array(s["GJ"]) / 1e6)
    ax[1].set_title("Torsion stiffness GJ (MN m²)", fontsize=9)
    ax[2].plot(y, s["m_span"])
    ax[2].set_title("Structural mass (kg/m)", fontsize=9)
    for a in ax:
        a.set_xlabel("y (m)", fontsize=8)
        a.grid(alpha=0.25)
        a.tick_params(labelsize=8)
    fig.tight_layout()
    fig.savefig(path, dpi=150)
    plt.close(fig)
    return path


def vg_figure(d, path):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    fig, ax = plt.subplots(1, 2, figsize=(9.5, 3.2))
    for fam, ls in (("symmetric", "-"), ("antisymmetric", "--")):
        f = d.get(fam) or {}
        V = f.get("V")
        g = f.get("g") or f.get("damping")
        fr = f.get("f") or f.get("freq")
        if not V or g is None:
            continue
        g, fr = np.array(g, float), np.array(fr, float) if fr is not None else None
        g = np.where(np.abs(g) > 2, np.nan, g)           # overdamped modes leave the plot
        for k in range(g.shape[1] if g.ndim == 2 else 0):
            ax[0].plot(V, g[:, k], ls, lw=1)
            if fr is not None and fr.ndim == 2:
                ax[1].plot(V, fr[:, k], ls, lw=1)
    ax[0].axhline(0, color="k", lw=0.8)
    ax[0].set_ylim(-1.0, 0.3)
    ax[0].set_title("Damping g (— symmetric, - - antisymmetric)", fontsize=9)
    ax[1].set_title("Frequency (Hz)", fontsize=9)
    for a in ax:
        a.set_xlabel("V (m/s)", fontsize=8)
        a.grid(alpha=0.25)
        a.tick_params(labelsize=8)
    if d.get("flutter_speed") and math.isfinite(d["flutter_speed"]):
        ax[0].axvline(d["flutter_speed"], color="#9C3D54", lw=1)
    fig.tight_layout()
    fig.savefig(path, dpi=150)
    plt.close(fig)
    return path


def change_figure(rows, variants, path):
    """horizontal bars: % change of every result against the first variant"""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    rr = [r for r in rows if any(r["change"].get(v) for v in variants[1:])]
    if not rr:
        return None
    labels = [f"{r['node']}: {r['metric']}" for r in rr]
    n, m = len(rr), len(variants) - 1
    h = 0.8 / max(m, 1)
    fig, ax = plt.subplots(figsize=(7.2, 0.28 * n * max(m, 1) + 1.2))
    pal = ["#1D4E89", "#C2571A", "#2B7A4B", "#7A4FA0", "#8A7B12"]
    for j, v in enumerate(variants[1:]):
        vals = [100 * (r["change"].get(v) or 0.0) for r in rr]
        ax.barh(np.arange(n) + j * h, vals, height=h, color=pal[j % len(pal)], label=v)
    ax.set_yticks(np.arange(n) + 0.4 - h / 2)
    ax.set_yticklabels(labels, fontsize=7)
    ax.invert_yaxis()
    ax.axvline(0, color="k", lw=0.8)
    ax.set_xlabel(f"change against {variants[0]} (%)", fontsize=8)
    ax.tick_params(axis="x", labelsize=8)
    ax.grid(axis="x", alpha=0.25)
    ax.legend(fontsize=8, loc="lower right")
    fig.tight_layout()
    fig.savefig(path, dpi=150)
    plt.close(fig)
    return path


# ----------------------------------------------------------------------
# content
# ----------------------------------------------------------------------
def _ac_of(p, variant, status):
    import aircraft as A
    for nid, s in status.items():
        if s["kind"] == "aircraft":
            f, _ = p.output(variant, nid, status=status)
            if f and f.exists():
                return A.Aircraft.load(f)
    return None


def _aircraft_section(ac, figdir):
    from weight_balance import pct_mac
    blocks = [H(2, "Aircraft")]
    rows = [["Mass", f"{ac.mass:.1f} kg"], ["Wing area", f"{ac.S:.3f} m²"], ["Span", f"{ac.b:.3f} m"],
            ["Aspect ratio", f"{ac.b ** 2 / ac.S:.2f}"], ["Mean aerodynamic chord", f"{ac.c:.3f} m"],
            ["Centre of gravity", f"x {ac.cg[0]:.3f} m ({pct_mac(ac, ac.cg[0]):.1f} % MAC)"]]
    blocks.append(TAB(["Quantity", "Value"], rows, widths=[2.5, 3.5]))
    srows = [[s.name, f"{s.span:.3f}", f"{s.root_chord:.3f}", f"{s.tip_chord:.3f}", f"{s.sweep:g}",
              f"{s.dihedral:g}", f"{s.twist:g}", s.airfoil + (f" → {s.tip_airfoil}" if s.tip_airfoil else "")]
             for s in ac.surfaces]
    blocks.append(TAB(["Surface", "span (m)", "root c (m)", "tip c (m)", "sweep (°)", "dihedral (°)",
                       "twist (°)", "section"], srows))
    blocks.append(IMG(planform_figure(ac, figdir / "planform.png"), "Planform (control surfaces shaded)", 5.2))
    return blocks


def _analysis_section(p, variant, nid, st, figdir):
    import project as PR
    import trust as T
    n = p.nodes(variant)[nid]
    kind = st["kind"]
    f, current = p.output(variant, nid, status=st and {nid: st})
    blocks = [H(2, n["label"])]
    if f is None or not f.exists():
        blocks.append(P(f"Not computed ({st['state']}).", "muted"))
        return blocks, None
    d = json.loads(f.read_text())
    if not current:
        blocks.append(P("This result is STALE: its inputs changed after it was computed ("
                        + "; ".join(st["reasons"]) + "). Recompute before relying on it.", "warn"))
    prm = {k: v for k, v in n["params"].items() if k != "overrides"}
    if prm:
        blocks.append(P("Inputs: " + ", ".join(f"{k} = {_fmt(v)}" for k, v in prm.items()), "muted"))
    mets = PR.METRICS.get(kind, [])
    rows = []
    for key, label, unit, better, get in mets:
        try:
            v = float(get(d))
        except Exception:                                # noqa: BLE001
            continue
        rows.append([label, _fmt(v), unit])
    if rows:
        blocks.append(TAB(["Result", "Value", "Unit"], rows, widths=[3.0, 1.6, 1.4]))
    # confidence
    if kind == "polar":
        ass = T.assess_polar(d)
        lev = ass["level"]
        blocks.append(P(f"Confidence: {T.summary_line(ass)}.", lev))
        bad = [x for x in ass["points"] if x["level"] != "green"]
        if bad:
            blocks.append(TAB(["α (°)", "level", "Cl ±", "Cd ±", "reason"],
                              [[_fmt(x["alpha"]), T.LABEL[x["level"]], f"{100 * x['cl_band']:.1f} %",
                                f"{100 * x['cd_band']:.1f} %",
                                next((t for l, t in x["reasons"] if l == x["level"]), "")[:160]] for x in bad],
                              colors={(i, 1): x["level"] for i, x in enumerate(bad)}, widths=[0.6, 1.1, 0.7, 0.7, 3.6]))
    else:
        ctx = {}
        if kind == "performance":
            ctx["flap_source"] = str((d.get("landing") or {}).get("flap_source", ""))
        if kind == "structure":
            fl = (d.get("aeroelastic") or {}).get("flutter") or {}
            vd = (d.get("aeroelastic") or {}).get("V_D")
            if fl.get("V_flutter") and vd:
                ctx["flutter_margin"] = fl["V_flutter"] / (1.2 * vd)
        notes = T.tool_notes(kind, **ctx)
        lev = T.worst(*(l for l, _ in notes)) if notes else None
        if notes:
            blocks.append(P(f"Confidence: {T.LABEL.get(lev, '')}.", lev))
            blocks.append(BUL(f"[{T.LABEL[l]}] {t}" for l, t in notes))
    # figures
    for fn in FIGURES.get(kind, []):
        if (f.parent / fn).exists():
            blocks.append(IMG(f.parent / fn, fn.replace("_", " ").replace(".png", ""), 6.3))
    try:
        if kind == "structure":
            blocks.append(IMG(structure_figure(d, figdir / f"{nid}_structure.png"), "Wing box along the span"))
        if kind == "aeroelastic":
            blocks.append(IMG(vg_figure(d, figdir / f"{nid}_vg.png"), "p-k flutter: damping and frequency"))
    except Exception:                                    # noqa: BLE001  (figure is optional)
        pass
    return blocks, lev


def design_report(p, variant=None, figdir=None):
    """blocks of the design report of one variant"""
    import trust as T
    variant = variant or p.baseline
    figdir = Path(figdir or (p.dir / "reports" / "_fig"))
    figdir.mkdir(parents=True, exist_ok=True)
    status = p.status(variant)
    vd = p.variants[variant]
    tpl = p.data.get("template") or {}
    blocks = [H(0, f"{p.name}: design report"),
              P(f"Variant: {variant}" + (" (baseline)" if variant == p.baseline else "")
                + (f" — {vd.get('description')}" if vd.get("description") else ""), "sub"),
              P(f"{dt.date.today().isoformat()} · AeroLab {_git_version()}"
                + (f" · template: {tpl.get('title')}" if tpl else ""), "muted")]
    if p.data.get("description"):
        blocks.append(P(p.data["description"]))
    # status
    counts = {}
    for s in status.values():
        counts[s["state"]] = counts.get(s["state"], 0) + 1
    blocks.append(H(1, "Summary"))
    st_txt = ", ".join(f"{n} {k}" for k, n in counts.items())
    blocks.append(P(f"{len(status)} analyses: {st_txt}.", "ok" if counts.get("up to date") == len(status) else "warn"))
    mets = p.metrics(variant, status)
    rows, colors = [], {}
    import project as PR
    for nid, s in status.items():
        m = mets.get(nid)
        if not m:
            continue
        for key, label, unit, better, _ in PR.METRICS.get(s["kind"], []):
            v = m["values"].get(key)
            if v is None:
                continue
            rows.append([s["label"], label, _fmt(v), unit, (T.LABEL.get(m["trust"], "") if m["trust"] else "")
                         + ("" if m["current"] else " (stale)")])
            if m["trust"]:
                colors[(len(rows) - 1, 4)] = m["trust"] if m["current"] else "amber"
    if rows:
        blocks.append(TAB(["Analysis", "Result", "Value", "Unit", "Confidence"], rows, colors,
                          widths=[1.5, 1.9, 1.0, 0.8, 1.6]))
    if p.diff(p.baseline, variant) and variant != p.baseline:
        blocks.append(P("Changes against the baseline: " + "; ".join(
            f"{n}.{k}: {_fmt(a)} → {_fmt(b)}" for n, k, a, b in p.diff(p.baseline, variant))))
    try:
        import preflight as PF
        pre = PF.check_project(p, variant)
        warn = [f"{v['label']}: {i['check']} — {i['detail']}" for v in pre["nodes"].values() for i in v["items"]
                if i["level"] in ("warn", "error")]
        if warn:
            blocks.append(H(2, "Setup checks"))
            blocks.append(P("The pre-run checks flag these points; they affect how far the results can be used.",
                            "warn"))
            blocks.append(BUL(warn))
    except Exception:                                    # noqa: BLE001 (checks are advisory here)
        pass
    if tpl.get("notes"):
        blocks.append(H(2, "Notes from the template"))
        blocks.append(BUL(tpl["notes"]))
    ac = _ac_of(p, variant, status)
    blocks.append(BREAK())
    if ac is not None:
        blocks += _aircraft_section(ac, figdir)
    used_kinds = set()
    for nid, s in status.items():
        if s["kind"] == "aircraft":
            continue
        bl, _ = _analysis_section(p, variant, nid, s, figdir)
        blocks += bl
        used_kinds.add(s["kind"])
    blocks += _assumptions_and_refs(used_kinds, p, variant, status)
    return blocks


def _assumptions_and_refs(kinds, p, variant, status):
    import trust as T
    blocks = [BREAK(), H(1, "Assumptions, limits and validation")]
    rows = []
    for k in sorted(kinds):
        ev = T.EVIDENCE.get(k)
        if ev:
            rows.append([k.replace("_", " ").capitalize(), T.LABEL[ev[0]], ev[1], ev[2]])
    if rows:
        blocks.append(TAB(["Tool", "Level", "Evidence and limits", "Report"], rows,
                          colors={(i, 1): T.EVIDENCE[k][0] for i, k in enumerate(sorted(x for x in kinds
                                                                                         if x in T.EVIDENCE))},
                          widths=[1.2, 1.0, 3.6, 1.2]))
    if "polar" in kinds:
        blocks.append(H(2, "Validated cases behind the section results"))
        blocks.append(TAB(["Case", "Reference", "Re", "Cl / Cd error"],
                          [[c.title, c.reference, f"{c.re[0]:.2g}–{c.re[1]:.2g}",
                            f"{100 * c.cl_err:.1f} % / {100 * c.cd_err:.1f} %"] for c in T.CASES],
                          widths=[2.8, 2.2, 1.0, 1.0]))
    blocks.append(H(2, "Traceability"))
    blocks.append(TAB(["Analysis", "State", "Result id", "Inputs"],
                      [[s["label"], s["state"], s["dir"].name if s["state"] == "up to date" else "—",
                        ", ".join(f"{k}={_fmt(v)}" for k, v in p.nodes(variant)[nid]["params"].items()
                                  if k != "overrides")[:180]]
                       for nid, s in status.items()], widths=[1.4, 0.9, 1.9, 2.8]))
    blocks.append(P("Every result is stored under the fingerprint of its inputs (Result id): rerunning a "
                    "project with the same inputs reproduces it. Levels: validated range = inside a case where "
                    "the code matched the reference; use with care = near one, or a known weak spot; low "
                    "confidence = outside every validated case. Reports in validation/ of the repository.",
                    "muted"))
    return blocks


def trade_report(p, variants=None, figdir=None):
    """blocks of a trade study between variants (first = reference)"""
    import project as PR
    import trust as T
    variants = variants or [p.baseline] + [v for v in p.variants if v != p.baseline]
    figdir = Path(figdir or (p.dir / "reports" / "_fig"))
    figdir.mkdir(parents=True, exist_ok=True)
    ref = variants[0]
    rows = p.compare(variants)
    blocks = [H(0, f"{p.name}: trade study"),
              P(" vs ".join(variants) + f" (reference: {ref})", "sub"),
              P(f"{dt.date.today().isoformat()} · AeroLab {_git_version()}", "muted")]
    blocks.append(H(1, "Variants"))
    vrows = []
    for v in variants:
        d = p.diff(ref, v) if v != ref else []
        vrows.append([v, p.variants[v].get("description", ""),
                      "reference" if v == ref else ("; ".join(f"{n}.{k}: {_fmt(a)} → {_fmt(b)}" for n, k, a, b in d)
                                                    or "no parameter differences")])
    blocks.append(TAB(["Variant", "Description", f"Changes against {ref}"], vrows, widths=[1.2, 1.8, 4.0]))
    # findings
    findings, caveats = [], []
    for v in variants[1:]:
        scored = [(r, r["change"].get(v)) for r in rows if r["change"].get(v) is not None and r["better"]]
        scored = [(r, c) for r, c in scored if abs(c) >= 0.01]
        good = sorted([x for x in scored if (x[1] > 0) == (x[0]["better"] == "higher")], key=lambda x: -abs(x[1]))
        bad = sorted([x for x in scored if (x[1] > 0) != (x[0]["better"] == "higher")], key=lambda x: -abs(x[1]))
        if good:
            findings.append(f"{v}: better — " + ", ".join(f"{r['metric']} {100 * c:+.1f} %" for r, c in good[:4]))
        if bad:
            findings.append(f"{v}: worse — " + ", ".join(f"{r['metric']} {100 * c:+.1f} %" for r, c in bad[:4]))
        if not good and not bad:
            findings.append(f"{v}: no result changes by more than 1 % in a direction that matters.")
        st = [r for r in rows if r["current"].get(v) is False]
        if st:
            caveats.append(f"{v}: {len(st)} results are stale (recompute before deciding).")
        low = {r["node"] for r in rows if r["trust"].get(v) in ("amber", "red")}
        if low:
            caveats.append(f"{v}: results to use with care come from {', '.join(sorted(low))}.")
    blocks.append(H(1, "Key findings"))
    blocks.append(BUL(findings or ["No differences yet: run the variants."]))
    if caveats:
        blocks.append(P("Caveats:", "warn"))
        blocks.append(BUL(caveats))
    blocks.append(H(1, "Results"))
    trows, colors = [], {}
    for i, r in enumerate(rows):
        row = [r["node"], r["metric"] + (f" ({r['unit']})" if r["unit"] else "")]
        for j, v in enumerate(variants):
            x = r["values"][v]
            c = r["change"].get(v)
            row.append("—" if x is None else _fmt(x) + (f" ({100 * c:+.1f} %)" if c and v != ref else "")
                       + ("" if r["current"].get(v) in (True, None) else " *"))
            if c and v != ref and r["better"] and abs(c) >= 0.005:
                colors[(i, 2 + j)] = "green" if (c > 0) == (r["better"] == "higher") else "red"
        trows.append(row)
    blocks.append(TAB(["Analysis", "Result"] + variants, trows, colors))
    blocks.append(P("Green: better, red: worse (where a better direction exists). * = stale result.", "muted"))
    fig = change_figure(rows, variants, figdir / "trade_change.png")
    if fig:
        blocks.append(IMG(fig, f"Change of every result against {ref}", 6.3))
    kinds = {p.nodes(ref)[nid]["kind"] for nid in p.nodes(ref)}
    blocks += _assumptions_and_refs(kinds - {"aircraft"}, p, ref, p.status(ref))
    return blocks


# ----------------------------------------------------------------------
# renderers
# ----------------------------------------------------------------------
def to_docx(blocks, path):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10)
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Inches(0.9)

    def rgb(h):
        h = h.lstrip("#")
        return RGBColor(int(h[:2], 16), int(h[2:4], 16), int(h[4:], 16))
    for b in blocks:
        t = b[0]
        if t == "h":
            doc.add_heading(b[2], level=b[1])
        elif t == "p":
            para = doc.add_paragraph()
            run = para.add_run(b[1])
            style = b[2]
            if style in COLOR or style in ("warn", "ok"):
                run.font.color.rgb = rgb(COLOR.get({"warn": "amber", "ok": "green"}.get(style, style), "#56606C"))
                if style in ("warn", "red", "amber"):
                    run.bold = True
            if style == "muted":
                run.font.color.rgb = rgb("#56606C")
                run.font.size = Pt(9)
            if style == "sub":
                run.font.size = Pt(12)
        elif t == "bullets":
            for x in b[1]:
                doc.add_paragraph(str(x), style="List Bullet")
        elif t == "table":
            _, header, rows, colors, widths = b
            tb = doc.add_table(rows=1, cols=len(header))
            tb.style = "Light Grid Accent 1"
            for j, h in enumerate(header):
                c = tb.rows[0].cells[j]
                c.text = str(h)
                for r_ in c.paragraphs[0].runs:
                    r_.bold = True
                    r_.font.size = Pt(8.5)
            for i, row in enumerate(rows):
                cells = tb.add_row().cells
                for j, v in enumerate(row):
                    cells[j].text = str(v)
                    for r_ in cells[j].paragraphs[0].runs:
                        r_.font.size = Pt(8.5)
                        if (i, j) in colors:
                            r_.font.color.rgb = rgb(COLOR.get(colors[(i, j)], "#56606C"))
                            r_.bold = True
            if widths:
                for j, w in enumerate(widths):
                    for row in tb.rows:
                        row.cells[j].width = Inches(w)
            doc.add_paragraph()
        elif t == "img":
            _, pth, cap, width = b
            if Path(pth).exists():
                doc.add_picture(pth, width=Inches(min(width, 6.6)))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                if cap:
                    cp = doc.add_paragraph(cap)
                    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for r_ in cp.runs:
                        r_.italic = True
                        r_.font.size = Pt(8.5)
        elif t == "break":
            doc.add_page_break()
    _ = (OxmlElement, qn)
    doc.save(path)
    return path


def to_pdf(blocks, path):
    from reportlab.lib import colors as rc
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import (Image, KeepTogether, ListFlowable, ListItem, PageBreak, Paragraph,
                                    SimpleDocTemplate, Spacer, Table, TableStyle)
    from xml.sax.saxutils import escape
    ss = getSampleStyleSheet()
    base = ParagraphStyle("b", parent=ss["Normal"], fontName="Helvetica", fontSize=9.5, leading=12.5)
    cell = ParagraphStyle("c", parent=base, fontSize=7.8, leading=9.5)
    cellb = ParagraphStyle("cb", parent=cell, fontName="Helvetica-Bold")
    hs = {0: ParagraphStyle("h0", parent=base, fontName="Helvetica-Bold", fontSize=18, leading=22, spaceAfter=4),
          1: ParagraphStyle("h1", parent=base, fontName="Helvetica-Bold", fontSize=13.5, leading=17, spaceBefore=10,
                            spaceAfter=4, textColor=rc.HexColor("#1D4E89")),
          2: ParagraphStyle("h2", parent=base, fontName="Helvetica-Bold", fontSize=11, leading=14, spaceBefore=8,
                            spaceAfter=3)}
    cap = ParagraphStyle("cap", parent=base, fontSize=8, alignment=TA_CENTER, textColor=rc.HexColor("#56606C"))
    story = []
    W = A4[0] - 1.6 * inch
    for b in blocks:
        t = b[0]
        if t == "h":
            story.append(Paragraph(escape(b[2]), hs[min(b[1], 2)]))
        elif t == "p":
            style = b[2]
            col = {"warn": "amber", "ok": "green"}.get(style, style)
            st = base
            if style == "muted":
                st = ParagraphStyle("m", parent=base, fontSize=8.5, textColor=rc.HexColor("#56606C"))
            elif style == "sub":
                st = ParagraphStyle("s", parent=base, fontSize=12, leading=15)
            elif col in COLOR:
                st = ParagraphStyle("k", parent=base, textColor=rc.HexColor(COLOR[col]),
                                    fontName="Helvetica-Bold" if col in ("amber", "red") else "Helvetica")
            story.append(Paragraph(escape(b[1]), st))
            story.append(Spacer(1, 3))
        elif t == "bullets":
            story.append(ListFlowable([ListItem(Paragraph(escape(str(x)), base), leftIndent=10) for x in b[1]],
                                      bulletType="bullet", start="•", leftIndent=12))
            story.append(Spacer(1, 4))
        elif t == "table":
            _, header, rows, colors, widths = b
            data = [[Paragraph(escape(str(h)), cellb) for h in header]]
            for i, row in enumerate(rows):
                r = []
                for j, v in enumerate(row):
                    if (i, j) in colors:
                        r.append(Paragraph(f'<font color="{COLOR.get(colors[(i, j)], "#56606C")}"><b>'
                                           f"{escape(str(v))}</b></font>", cell))
                    else:
                        r.append(Paragraph(escape(str(v)), cell))
                data.append(r)
            if widths:
                tot = sum(widths)
                cw = [W * w / tot for w in widths]
            else:
                cw = [W / len(header)] * len(header)
            tb = Table(data, colWidths=cw, repeatRows=1)
            tb.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), rc.HexColor("#EEF1F4")),
                                    ("LINEBELOW", (0, 0), (-1, 0), 0.6, rc.HexColor("#17202B")),
                                    ("LINEBELOW", (0, 1), (-1, -1), 0.25, rc.HexColor("#CDD4DC")),
                                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                    ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]))
            story.append(tb)
            story.append(Spacer(1, 6))
        elif t == "img":
            _, pth, caption, width = b
            if Path(pth).exists():
                from reportlab.lib.utils import ImageReader
                iw, ih = ImageReader(pth).getSize()
                w = min(width * inch, W)
                h = w * ih / iw
                if h > 7.5 * inch:
                    h = 7.5 * inch
                    w = h * iw / ih
                story.append(KeepTogether([Image(pth, width=w, height=h)] +
                                          ([Paragraph(escape(caption), cap)] if caption else [])))
                story.append(Spacer(1, 6))
        elif t == "break":
            story.append(PageBreak())

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(rc.HexColor("#56606C"))
        canvas.drawString(0.8 * inch, 0.5 * inch, "AeroLab report")
        canvas.drawRightString(A4[0] - 0.8 * inch, 0.5 * inch, f"page {doc.page}")
        canvas.restoreState()
    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=0.8 * inch, rightMargin=0.8 * inch,
                            topMargin=0.75 * inch, bottomMargin=0.8 * inch, title=_title(blocks), author="AeroLab")
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return path


def _title(blocks):
    return next((b[2] for b in blocks if b[0] == "h" and b[1] == 0), "AeroLab report")


def write(blocks, path, fmt=None):
    path = Path(path)
    fmt = fmt or path.suffix.lstrip(".")
    path.parent.mkdir(parents=True, exist_ok=True)
    return to_docx(blocks, path) if fmt == "docx" else to_pdf(blocks, path)


def main(argv=None):
    import project as PR
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    d = sub.add_parser("design")
    d.add_argument("project")
    d.add_argument("--variant", default=None)
    d.add_argument("--format", default="pdf", choices=["pdf", "docx"])
    d.add_argument("--out", default=None)
    t = sub.add_parser("trade")
    t.add_argument("project")
    t.add_argument("--variants", default=None)
    t.add_argument("--format", default="pdf", choices=["pdf", "docx"])
    t.add_argument("--out", default=None)
    a = ap.parse_args(argv)
    p = PR.Project.load(a.project)
    stamp = dt.datetime.now().strftime("%Y%m%d-%H%M")
    if a.cmd == "design":
        v = a.variant or p.baseline
        out = Path(a.out or p.dir / "reports" / f"design_{PR.slugify(v)}_{stamp}.{a.format}")
        write(design_report(p, v), out, a.format)
    else:
        vs = a.variants.split(",") if a.variants else None
        out = Path(a.out or p.dir / "reports" / f"trade_{stamp}.{a.format}")
        write(trade_report(p, vs), out, a.format)
    print(f"[report] -> {out}")


if __name__ == "__main__":
    main()
